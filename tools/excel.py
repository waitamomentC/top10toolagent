from __future__ import annotations

import json
import os
import re
from pathlib import Path

from models.schemas import ToolResult
from tools.base import BaseTool

# ── 文件格式白名单 ────────────────────────────────────────────────────

EXCEL_EXTENSIONS = {".xlsx", ".xls"}
FORBIDDEN_EXTENSIONS = {
    ".csv", ".tsv", ".txt", ".json", ".xml", ".yaml", ".yml",
    ".pdf", ".doc", ".docx", ".ppt", ".pptx",
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".svg",
    ".html", ".htm", ".py", ".js", ".ts", ".java", ".go",
}
FILE_EXT_RE = re.compile(r"\.(\w+)\b", re.I)


def validate_excel_path(file_path: str) -> str | None:
    """校验文件扩展名是否为 Excel 格式，返回错误信息或 None"""
    ext = Path(file_path).suffix.lower()
    if not ext:
        return "未指定文件扩展名，仅支持 Microsoft Excel (.xlsx / .xls) 格式。"
    if ext not in EXCEL_EXTENSIONS:
        return f"不支持的文件格式 '{ext}'，仅支持 Microsoft Excel (.xlsx / .xls) 格式。"
    return None


def detect_forbidden_format(query: str) -> str | None:
    """检测 query 中是否包含非 Excel 文件扩展名，用于路由层拦截"""
    exts = {"." + m.group(1).lower() for m in FILE_EXT_RE.finditer(query)}
    bad = exts & FORBIDDEN_EXTENSIONS
    if bad:
        return f"不支持的文件格式: {', '.join(sorted(bad))}。仅支持 Microsoft Excel (.xlsx / .xls) 文件操作。"
    return None


# ── 工具实现 ──────────────────────────────────────────────────────────


class ReadExcelTool(BaseTool):
    name = "read_excel"
    description = (
        "读取本地 Excel 文件（仅支持 .xlsx / .xls）。"
        "输入为文件路径，返回表格内容（前 100 行）。"
    )

    async def execute(self, input_str: str) -> ToolResult:
        file_path = input_str.strip().strip("\"'")

        err = validate_excel_path(file_path)
        if err:
            return ToolResult(success=False, data="", error=err)

        if not os.path.exists(file_path):
            return ToolResult(success=False, data="", error=f"文件不存在: {file_path}")

        try:
            ext = Path(file_path).suffix.lower()
            if ext == ".xlsx":
                result = self._read_xlsx(file_path)
            else:
                result = self._read_xls(file_path)
            return ToolResult(success=True, data=result)
        except Exception as e:
            return ToolResult(success=False, data="", error=str(e))

    def _read_xlsx(self, path: str) -> str:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts: list[str] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"--- Sheet: {sheet_name} ---")
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                if row_count >= 100:
                    parts.append("... (超过 100 行，已截断)")
                    break
                parts.append("\t".join(str(c) if c is not None else "" for c in row))
                row_count += 1
        wb.close()
        return "\n".join(parts)

    def _read_xls(self, path: str) -> str:
        import xlrd

        wb = xlrd.open_workbook(path)
        parts: list[str] = []
        for sheet_name in wb.sheet_names():
            ws = wb.sheet_by_name(sheet_name)
            parts.append(f"--- Sheet: {sheet_name} ---")
            for i in range(min(ws.nrows, 100)):
                parts.append("\t".join(str(ws.cell_value(i, j)) for j in range(ws.ncols)))
            if ws.nrows > 100:
                parts.append("... (超过 100 行，已截断)")
        return "\n".join(parts)


class WriteExcelTool(BaseTool):
    name = "write_excel"
    description = (
        "写入本地 Excel 文件（仅支持 .xlsx）。"
        "输入为 JSON: {\"file\": \"路径.xlsx\", \"sheet\": \"Sheet名\", \"data\": [[\"列1\",\"列2\"],[1,2]]}。"
        "data 第一行为表头，后续行为数据行。"
    )

    async def execute(self, input_str: str) -> ToolResult:
        try:
            params = json.loads(input_str.strip())
        except json.JSONDecodeError as e:
            return ToolResult(success=False, data="", error=f"输入不是有效 JSON: {e}")

        file_path = params.get("file", "")
        sheet_name = params.get("sheet", "Sheet1")
        data = params.get("data", [])

        if not file_path:
            return ToolResult(success=False, data="", error="缺少 'file' 参数")
        if not data or not isinstance(data, list):
            return ToolResult(success=False, data="", error="'data' 必须是非空二维数组")

        err = validate_excel_path(file_path)
        if err:
            return ToolResult(success=False, data="", error=err)

        # 仅支持写入 xlsx
        if Path(file_path).suffix.lower() != ".xlsx":
            return ToolResult(success=False, data="", error="写入仅支持 .xlsx 格式")

        try:
            import openpyxl

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name

            for row in data:
                ws.append(row)

            wb.save(file_path)
            wb.close()
            return ToolResult(
                success=True,
                data=f"成功写入 {file_path}，共 {len(data)} 行（含表头）。",
            )
        except Exception as e:
            return ToolResult(success=False, data="", error=str(e))
